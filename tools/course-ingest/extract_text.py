"""Extract plain text from downloaded PDF / DOCX / HTML / RTF / XLSX into course/text/<chat>/<msg>.txt (idempotent)."""
import json, re, html, sys
from pathlib import Path
from common import COURSE

def pdf(p):
    import fitz
    with fitz.open(p) as d:
        return "\n\n".join(f"[page {i+1}]\n" + pg.get_text() for i, pg in enumerate(d))
def docx(p):
    import docx as D
    d = D.Document(p)
    out = [para.text for para in d.paragraphs]
    for t in d.tables:
        for row in t.rows: out.append(" | ".join(c.text for c in row.cells))
    return "\n".join(out)
def htm(p):
    s = Path(p).read_text(encoding="utf-8", errors="replace")
    s = re.sub(r"(?is)<(script|style).*?</\1>", "", s)
    s = re.sub(r"(?i)<br\s*/?>|</p>|</div>|</h\d>|</li>", "\n", s)
    return html.unescape(re.sub(r"<[^>]+>", "", s))
def rtf(p):
    s = Path(p).read_text(errors="replace")
    return re.sub(r"\[a-z]+-?\d* ?|[{}]", "", s)
def xlsx(p):
    import zipfile
    z = zipfile.ZipFile(p)
    ss = z.read("xl/sharedStrings.xml").decode("utf8", "replace") if "xl/sharedStrings.xml" in z.namelist() else ""
    return "\n".join(html.unescape(x) for x in re.findall(r"<t[^>]*>(.*?)</t>", ss, flags=re.S))

HANDLERS = {".pdf": pdf, ".docx": docx, ".html": htm, ".htm": htm, ".rtf": rtf, ".xlsx": xlsx}

def main():
    (COURSE / "text").mkdir(exist_ok=True)
    n = bad = 0
    seen = set()
    idx = COURSE / "index.jsonl"
    for l in idx.read_text(encoding="utf-8").splitlines():
        r = json.loads(l)
        if r.get("dup") or r["path"] in seen: continue
        seen.add(r["path"])
        ext = Path(r["path"]).suffix.lower()
        if ext not in HANDLERS: continue
        out = COURSE / "text" / r["chat"] / f"{r['msg']}.txt"
        if out.exists(): continue
        try:
            txt = HANDLERS[ext](COURSE / r["path"])
            out.parent.mkdir(parents=True, exist_ok=True)
            out.write_text(txt, encoding="utf-8"); n += 1
        except Exception as e:
            bad += 1; print("FAIL", r["path"], type(e).__name__, str(e)[:80])
    print("extracted", n, "failed", bad)
if __name__ == "__main__": main()
